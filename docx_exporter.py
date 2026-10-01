"""
docx_exporter.py - DOCX Document Generator for LegalEase
Generates professional, styled Microsoft Word (.docx) documents from legal text.
"""

import io
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def create_docx_from_text(title: str, text_content: str, language: str = "en") -> io.BytesIO:
    """
    Creates a styled .docx document from the legal text.
    Returns an in-memory BytesIO buffer ready for Streamlit / FastAPI download.
    """
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Document Header
    header = doc.sections[0].header
    hp = header.paragraphs[0]
    hp.text = "LegalEase AI-Generated Legal Document"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if hp.runs:
        hp.runs[0].font.size = Pt(8.5)
        hp.runs[0].font.color.rgb = RGBColor(120, 120, 120)

    # Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(title.upper())
    title_run.font.size = Pt(16)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(15, 23, 42)
    title_p.paragraph_format.space_after = Pt(12)

    lines = text_content.split("\n")
    skip_header = False

    for line in lines:
        stripped = line.strip()

        if stripped.startswith("===") or stripped.startswith("---"):
            continue

        if stripped.lower() == title.lower() and not skip_header:
            skip_header = True
            continue

        p = doc.add_paragraph()

        is_heading = (
            (stripped.isupper() and len(stripped) > 3 and not stripped.startswith("_"))
            or stripped.startswith("விதிமுறைகள்")
            or stripped.startswith("முக்கிய விதிமுறைகள்")
            or stripped.startswith("TERMS AND CONDITIONS")
            or stripped.startswith("SIGNATURES:")
            or stripped.startswith("கையொப்பங்கள்:")
        )

        if is_heading:
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(stripped)
            run.font.bold = True
            run.font.size = Pt(11.5)
            run.font.color.rgb = RGBColor(30, 41, 59)
        elif "____" in stripped:
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(stripped)
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(70, 70, 70)
        else:
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(stripped)
            run.font.size = Pt(10.5)
            run.font.color.rgb = RGBColor(40, 40, 40)

    # Footer
    footer = doc.sections[0].footer
    fp = footer.paragraphs[0]
    fp.text = "Generated via LegalEase | Please review with a legal professional."
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if fp.runs:
        fp.runs[0].font.size = Pt(8)
        fp.runs[0].font.color.rgb = RGBColor(130, 130, 130)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
